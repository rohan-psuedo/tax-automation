"""Word and spreadsheet files. These have no page images: their content is text or rows."""

import csv
import io
import zipfile
from pathlib import Path

import docx
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from app.parsing.base import ParsedDocument, ParsedSheet, ParseError

MAX_SHEET_ROWS = 500


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop fully empty rows and trailing empty columns."""
    rows = [r for r in rows if any(r)]
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    return [r[:width] + [""] * (width - len(r[:width])) for r in rows]


def _sheet_text(sheets: list[ParsedSheet]) -> str:
    parts = []
    for sheet in sheets:
        lines = "\n".join(" | ".join(row) for row in sheet.rows)
        parts.append(f"--- Sheet {sheet.name} ---\n{lines}")
    return "\n\n".join(parts)


def parse_docx(path: Path) -> ParsedDocument:
    try:
        document = docx.Document(str(path))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ParseError("This Word file is damaged and can't be opened.") from exc

    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tables = []
    for t_index, table in enumerate(document.tables, start=1):
        rows = _trim([[cell.text.strip() for cell in row.cells] for row in table.rows])
        if rows:
            tables.append(ParsedSheet(name=f"Table {t_index}", rows=rows, total_rows=len(rows)))
            blocks.append("\n".join(" | ".join(r) for r in rows))

    text = "\n\n".join(blocks)
    if not text:
        raise ParseError("This Word file has no text in it.")
    return ParsedDocument(text=text, has_text_layer=True, sheets=tables)


def parse_xlsx(path: Path) -> ParsedDocument:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ParseError("This Excel file is damaged and can't be opened.") from exc

    result = ParsedDocument(has_text_layer=True)
    try:
        for ws in workbook.worksheets:
            rows: list[list[str]] = []
            total = 0
            for raw in ws.iter_rows(values_only=True):
                if not any(v is not None and str(v).strip() for v in raw):
                    continue
                total += 1
                if len(rows) < MAX_SHEET_ROWS:
                    rows.append([_cell(v) for v in raw])
            if rows:
                result.sheets.append(ParsedSheet(name=ws.title, rows=_trim(rows), total_rows=total))
                if total > MAX_SHEET_ROWS:
                    result.warnings.append(
                        f"Sheet '{ws.title}' has {total} rows; only the first {MAX_SHEET_ROWS} "
                        "are shown."
                    )
    finally:
        workbook.close()

    if not result.sheets:
        raise ParseError("This Excel file has no data in it.")
    result.text = _sheet_text(result.sheets)
    return result


def parse_csv(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            content = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes nearly anything
        raise ParseError("This CSV file uses an unknown text encoding.")

    try:
        dialect = csv.Sniffer().sniff(content[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    all_rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(content), dialect)]
    all_rows = [r for r in all_rows if any(r)]
    if not all_rows:
        raise ParseError("This CSV file is empty.")

    sheet = ParsedSheet(
        name=path.stem, rows=_trim(all_rows[:MAX_SHEET_ROWS]), total_rows=len(all_rows)
    )
    result = ParsedDocument(has_text_layer=True, sheets=[sheet])
    if len(all_rows) > MAX_SHEET_ROWS:
        result.warnings.append(
            f"This file has {len(all_rows)} rows; only the first {MAX_SHEET_ROWS} are shown."
        )
    result.text = _sheet_text(result.sheets)
    return result

"""Builds small sample documents in memory, so tests need no binary fixtures."""

import io

import docx
import openpyxl
import pymupdf
from PIL import Image, ImageDraw

INVOICE_TEXT = [
    "TAX INVOICE",
    "Sharma Electronics, Bengaluru  GSTIN 29ABCDE1234F1ZW",
    "Invoice No: SE/2026/0042   Date: 28-09-2026",
    "Network switch 24-port   Qty 2   Rate 5,000.00   Amount 10,000.00",
    "CGST 9% 900.00   SGST 9% 900.00   Total 11,800.00",
]


def text_pdf(lines: list[str] = INVOICE_TEXT, pages: int = 1) -> bytes:
    doc = pymupdf.open()
    for n in range(pages):
        page = doc.new_page()
        for i, line in enumerate(lines):
            page.insert_text((72, 72 + i * 18), line if n == 0 else f"{line} (page {n + 1})")
    return doc.tobytes()


def png(size: tuple[int, int] = (1200, 1600), text: str = "TAX INVOICE") -> bytes:
    img = Image.new("RGB", size, "white")
    ImageDraw.Draw(img).text((50, 50), text, fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def scanned_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=png())
    return doc.tobytes()


def password_pdf() -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")


def rotated_jpeg() -> bytes:
    """A landscape-stored photo whose EXIF says 'rotate 90°' (like a phone camera)."""
    img = Image.new("RGB", (1600, 1200), "white")
    exif = Image.Exif()
    exif[0x0112] = 6  # Orientation: rotate 90 CW to display
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif)
    return buf.getvalue()


def docx_file() -> bytes:
    document = docx.Document()
    for line in INVOICE_TEXT[:3]:
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=3)
    for col, value in enumerate(["Item", "Qty", "Amount"]):
        table.cell(0, col).text = value
    for col, value in enumerate(["Network switch", "2", "10000"]):
        table.cell(1, col).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def xlsx_file(rows: int = 3) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Invoice", "Party", "Amount"])
    for i in range(rows):
        ws.append([f"INV-{i + 1}", "Gupta Retail", 1000.0 * (i + 1)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_file() -> bytes:
    return b"Date;Narration;Debit;Credit\n01-09-2026;NEFT Sharma Electronics;11800;\n"
